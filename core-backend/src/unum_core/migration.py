"""Stream CSV/XLSX into SQLite, PostgreSQL or MySQL with inferred types."""
from __future__ import annotations

import asyncio
import re
import unicodedata
from collections.abc import Awaitable, Callable, Iterator
from datetime import date, datetime
from itertools import islice
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncEngine

from .database import Database, identifier
from .env_manager import EnvironmentManager
from .relational import quote, render_type

CHUNK_SIZE = 1000
TYPE_ORDER = {"INTEGER": 0, "REAL": 1, "TEXT": 2}


def normal_name(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", str(value).strip()).encode("ascii", "ignore").decode("ascii")
    name = re.sub(r"[^A-Za-z0-9_]+", "_", normalized).strip("_").lower()
    if not name or name[0].isdigit():
        name = "col_" + name
    return name


def column_names(raw: list[object]) -> list[str]:
    names: list[str] = []
    for item in raw:
        base = normal_name(item)
        name = base
        index = 2
        while name in names:
            name = f"{base}_{index}"
            index += 1
        names.append(name)
    if not names or len(names) > 500:
        raise ValueError("Import requires 1 to 500 columns")
    return names


def infer_type(series: pd.Series) -> str:
    present = series.dropna()
    if present.empty:
        return "INTEGER"
    if pd.api.types.is_integer_dtype(present) or pd.api.types.is_bool_dtype(present):
        return "INTEGER"
    if pd.api.types.is_float_dtype(present):
        return "REAL"
    if pd.api.types.is_datetime64_any_dtype(present):
        return "TEXT"
    if all(isinstance(value, (bool, int)) for value in present):
        return "INTEGER"
    if all(isinstance(value, (bool, int, float)) for value in present):
        return "REAL"
    return "TEXT"


def value_for_sql(value: object) -> object:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, (date, datetime, pd.Timestamp)):
        return value.isoformat()
    if hasattr(value, "item"):
        return value.item()
    return value


class Migration:
    def __init__(self, env: EnvironmentManager, db: Database):
        self.env, self.db = env, db

    def _path(self, raw: str) -> Path:
        path = self.env.project_path(raw)
        if not path.is_file() or path.suffix.lower() not in {".csv", ".xlsx"}:
            raise ValueError("Select a CSV or XLSX file in the workspace")
        if path.stat().st_size > 100_000_000:
            raise ValueError("Import exceeds 100 MB")
        return path

    def _chunks(self, path: Path) -> Iterator[pd.DataFrame]:
        if path.suffix.lower() == ".csv":
            chunks = pd.read_csv(path, chunksize=CHUNK_SIZE)
            names: list[str] | None = None
            for chunk in chunks:
                if names is None:
                    names = column_names(list(chunk.columns))
                chunk.columns = names
                yield chunk
            return
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                raise ValueError("Spreadsheet is empty")
            names = column_names(list(header))
            while batch := list(islice(rows, CHUNK_SIZE)):
                yield pd.DataFrame(batch, columns=names)
        finally:
            workbook.close()

    def _preview(self, path: Path, name: str, dialect: str = "sqlite") -> dict:
        if dialect not in {"sqlite", "postgresql", "mysql"}:
            raise ValueError("Unsupported import target")
        columns: list[str] | None = None
        types: dict[str, str] = {}
        sample: list[list[str]] = []
        count = 0
        for chunk in self._chunks(path):
            columns = list(chunk.columns)
            count += len(chunk)
            for column in columns:
                inferred = infer_type(chunk[column])
                previous = types.get(column, "INTEGER")
                types[column] = max((previous, inferred), key=lambda kind: TYPE_ORDER[kind])
            if len(sample) < 5:
                sample.extend(chunk.head(5 - len(sample)).fillna("").astype(str).values.tolist())
        if not columns:
            raise ValueError("File has no data rows")
        definitions = [{"name": column, "type": types[column]} for column in columns]
        ident = identifier if dialect == "sqlite" else lambda value: quote(value, dialect)
        ddl = f"CREATE TABLE {ident(name)} (" + ", ".join(f"{ident(col['name'])} {col['type'] if dialect == 'sqlite' else render_type(col['type'], dialect)}" for col in definitions) + ");"
        return {"table": name, "columns": definitions, "rows": count, "sample": sample, "ddl": ddl}

    async def preview(self, raw: str, table: str | None = None, dialect: str = "sqlite") -> dict:
        path = self._path(raw)
        return await asyncio.to_thread(self._preview, path, normal_name(table or path.stem), dialect)

    async def import_file(self, raw: str, table: str | None = None, engine: AsyncEngine | None = None, dialect: str = "sqlite", progress: Callable[[dict], Awaitable[None]] | None = None) -> dict:
        path = self._path(raw)
        name = normal_name(table or path.stem)
        preview = await asyncio.to_thread(self._preview, path, name, dialect)
        names = [column["name"] for column in preview["columns"]]
        ident = identifier if dialect == "sqlite" else lambda value: quote(value, dialect)
        insert = f"INSERT INTO {ident(name)} (" + ", ".join(map(ident, names)) + ") VALUES (" + ", ".join(f":{column}" for column in names) + ")"
        inserted = 0
        async with (engine or self.db.engine).begin() as connection:
            await connection.execute(text(preview["ddl"]))
            chunks = self._chunks(path)
            while (chunk := await asyncio.to_thread(lambda: next(chunks, None))) is not None:
                records = await asyncio.to_thread(lambda: [{column: value_for_sql(value) for column, value in row.items()} for row in chunk.to_dict("records")])
                await connection.execute(text(insert), records)
                inserted += len(records)
                if progress:
                    await progress({"event": "IMPORT_PROGRESS", "inserted": inserted, "total": preview["rows"], "percent": round(inserted / preview["rows"] * 100, 1)})
        return {"table": name, "inserted": inserted}
